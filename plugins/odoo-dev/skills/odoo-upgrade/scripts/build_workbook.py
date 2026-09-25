#!/usr/bin/env python3
"""Merge inventory CSV + agent JSON into the upgrade-preparation workbook.

Inputs, all in one working dir (see references/inventory.md,
references/oca-base-review.md, references/functional-requirements.md):

    inventory.csv     seeded by module_inventory.py, OCA repo column filled by
                      oca_check.py
    enrich_*.json     phase 1 — per-module enrichment + "<major> native?"
    oca_alt_*.json    phase 2 — per-module "OCA <major> alternative"
    fr_*.json         phase 3 — functional requirements (optional)
    fr_merges.json    optional {drop_tmp_id: {keep, note, requirement}}
    fr_notes.json     optional {tmp_id: "note appended to Notes"}
    studio.csv        optional — studio_inventory.py output (UI-built artifacts)
    tickets.csv       optional — Odoo support register, see references/support-tickets.md

Both agent groups are merged FIELD BY FIELD, per module, across every file in
the group: an agent may deliver a partial record, and a second file naming the
same module fills the keys the first left out instead of replacing the record.
A key carrying different non-empty values in two files is a problem line naming
the module, the key and both files; the first file read (files sorted by name)
keeps the cell.

Output: one xlsx with Module Inventory / Inventory Evidence / Functional
Requirements / Traceability, plus Studio and Tickets when those CSVs are
present, plus CSV siblings. Everything else the reader
needs (column semantics, writing rules) lives in the skill, not in the file.

    python3 build_workbook.py --workdir /tmp/inv --target-version 19.0 \
            -o /mnt/extra-addons/acme_upgrade_16_to_19_workbook.xlsx

Prints a validation report — unenriched modules, conflicting enrichment
files, `TODO-AI` sentinels in a delivered column, over-long cells, inventory
cells outside their documented vocabulary (Functional area, native verdict, OCA
repo, OCA alternative), requirements with weak wording or no (case-insensitive)
"shall", a requirement whose `sources` is not a list of module names, unknown
source modules, modules with no requirement, and near-duplicate requirements
across agent groups. Read it: it is the only check that the fan-out agents
followed the briefs. `problems: 0` means complete AND schema-valid: a column
still carrying the sentinel is counted, reported and exits non-zero.
"""

import argparse
import csv
import itertools
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_COLUMNS = ["Module Name", "Display Name", "Module Purpose",
                "Source version", "LoC", "Dependencies", "Dependency origin",
                "3rd party app?", "3rd party app link", "OCA repo",
                "Complexity/risk", "Upgrade action", "Functional area"]
EVIDENCE_COLUMNS = ["Module", "native? evidence", "OCA alternative evidence",
                    "Vendor release", "Complexity rationale",
                    "Upgrade action rationale", "Notes"]
FR_KEYS = {"tmp_id", "requirement", "type", "functional_area", "actor",
           "sources", "evidence", "status_source", "status_note", "handled",
           "handled_by", "handled_notes", "verification", "notes"}
# The ONE closed functional-area vocabulary. Both halves of the workbook are
# validated against this list — the inventory's "Functional area" column and
# every requirement's functional_area — so the two cannot drift apart and an
# area copied from a module row into a requirement always passes.
AREA_ORDER = ["Sales", "CRM", "Purchasing", "Inventory", "Manufacturing",
              "Accounting", "HR", "Website/Portal", "Reporting", "Integration",
              "Technical/Base"]
# "OCA repo": verified org location, or the negative literal. Anything else is
# provenance prose — it belongs in the Evidence sheet, not a client-facing cell.
# The two literals are written by module_inventory.oca_seed (same names there)
# and by odoo-prior-art/scripts/oca_check.py; this is the reader of both.
OCA_NONE = "none"
OCA_UNVERIFIED = "claimed — run oca_check.py"
OCA_REPO_RE = re.compile(r"^OCA/[A-Za-z0-9._-]+(?: \([^()]*\))?$")
# "OCA <major> alternative": none, an already-OCA statement for a module that
# IS the upstream module, or up to three candidates.
OCA_ALREADY_RE = re.compile(
    r"^already OCA: [A-Za-z0-9._-]+/[A-Za-z0-9._-]+$")
OCA_ALT_RE = re.compile(
    r"^[A-Za-z0-9._-]+/[A-Za-z0-9._-]+ \((?:full|partial)\)$")
STATUSES = {"active", "dead", "broken"}
STUDIO_COLUMNS = ["kind", "technical_name", "model", "label", "owner_module",
                  "active", "classification", "detail"]
TICKET_COLUMNS = ["id", "date", "subject", "token", "status",
                  "blocking_module", "resolution", "link"]
STUDIO_CLASSIFICATIONS = {"purge", "convert-to-code", "keep-as-data", "review"}
TICKET_STATUSES = {"open", "waiting-odoo", "waiting-us", "resolved", "waived"}
# "<major> native?" is a VERDICT and nothing else: a closed vocabulary can be
# filtered and counted in a spreadsheet, which is the single biggest cost lever
# in the exercise. The explanation lives in "<major> native notes" beside it.
NATIVE_VERDICTS = ("yes", "yes/partial", "partial", "no")
NATIVE_NOTES_MAX = 300
# Every evidence cell is read inside one spreadsheet cell, and it is the
# reader's only audit trail back to the source tree; a runaway entry displaces
# the columns beside it. Same cap as references/inventory.md documents.
EVIDENCE_MAX = 300
# What the seeding script writes when only an agent can answer. In a column the
# workbook presents as an answer it is an unanswered question, not a value.
SENTINEL = "TODO-AI"
# A word inside `backticks` or "double quotes" is a literal the client uses —
# a delivery type called `normal`, a report called "100% invoiced" — so it is
# never weak wording. Masked out (length preserved) before the rules run.
QUOTED_LITERAL = re.compile(r'`[^`]*`|"[^"]*"')
# The weak-wording rules, one (label, pattern) per defect class, each searched
# as a predicate over the masked text instead of one flat substring grep: the
# label is what makes a hit actionable, and the shape of the pattern is what
# keeps `may` the modal apart from May the month. This list is the source of
# truth for rule 5 of references/functional-requirements.md — change both
# together, one doc line per label.
WEAK_WORDS = [
    # Modal instead of `shall`, recognised by the bare verb that follows it,
    # so "in May 2026" and "the will of the client" are left alone.
    ("modal verb",
     re.compile(r"(?i:\b(?:should|must|will|may|might|could)\b)"
                r"(?=\s+(?:not\s+)?(?!of\b|the\b|an?\b)[a-z]{2,}\b)")),
    # "shall be able to X" states a capability, not a behaviour: untestable.
    ("superfluous infinitive",
     re.compile(r"\bshall\s+(?:be\s+(?:able\s+to|capable\s+of)"
                r"|have\s+the\s+ability\s+to"
                r"|provide\s+the\s+ability\s+to|support|allow)\b"
                r"|\bis\s+designed\s+to\b", re.I)),
    # An open list cannot be accepted or rejected — the client signs a blank.
    ("open-ended clause",
     re.compile(r"\betc\.|\band so on\b|\bincluding but not limited to\b"
                r"|\bsuch as\b|\be\.g\.", re.I)),
    ("and/or", re.compile(r"\band\s*/\s*or\b", re.I)),
    ("absolute",
     re.compile(r"\b(?:always|never|all|every|none)\b|\b100%", re.I)),
    # A comparative needs its baseline: "faster than the 16.0 report".
    ("bare comparative",
     re.compile(r"\b(?:faster|better|more|less|improved|enhanced"
                r"|optimi[sz]ed)\b(?!\s+than\b)", re.I)),
    # A pronoun opening a clause has no antecedent a tester can resolve.
    ("pronoun without antecedent",
     re.compile(r"(?:^|[,;:]\s*)(?P<w>it|this|that|these|those|they)\b",
                re.I)),
    ("vague adjective",
     re.compile(r"\b(?:fast|easy|user-friendly|efficient|appropriate|normal"
                r"|few|most|timely|properly|quickly|reliable|intuitive"
                r"|seamless|robust|as needed|if possible)\b", re.I)),
]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
TOP_WRAP = Alignment(vertical="top", wrap_text=True)


def load_group_json(workdir, pattern):
    """Every entry in the group, tagged with the file it came from.

    The file name is carried because it is the only thing that makes a
    conflict actionable: "two files disagree" is useless without their names.
    """
    merged = []
    for path in sorted(workdir.glob(pattern)):
        for entry in json.loads(path.read_text()):
            merged.append((entry, path.name))
    return merged


def merge_group_json(workdir, pattern, problems):
    """Field-wise merge of one agent group, keyed by module.

    A fan-out writes one file per group, and nothing stops two groups from
    covering the same module — a record-wise merge would keep whichever file
    sorted last and silently drop every key the loser held (evidence keys have
    no seed CSV to fall back on, so they would simply ship empty).

    So: a key set by exactly one file wins; the same value in several files is
    fine; different non-empty values are a problem line naming the module, the
    key and both files, and the first file read keeps the cell. An empty value
    never overwrites a filled one.
    """
    merged, source = {}, {}
    for entry, name in load_group_json(workdir, pattern):
        module = (entry.get("module") or "").strip()
        if not module:
            problems.append(f"{name}: entry with no module key — dropped")
            continue
        record = merged.setdefault(module, {})
        origin = source.setdefault(module, {})
        for key, value in entry.items():
            if value in ("", None):
                continue
            if key not in record or record[key] in ("", None):
                record[key], origin[key] = value, name
            elif record[key] != value:
                problems.append(
                    f"{module}: {origin[key]} and {name} disagree on {key} "
                    f"— keeping {origin[key]}")
    return merged


def style_sheet(sheet, widths, freeze):
    for row in sheet.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = TOP_WRAP
    for cell in sheet[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions


def write_csv(path, columns, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(columns)
        writer.writerows(rows)


def build_passthrough_sheet(workbook, path, title, columns, widths, problems,
                            validate=None):
    """Fold a flat CSV in as its own sheet, validating what can be validated.

    Studio rows and support tickets are produced elsewhere (studio_inventory.py,
    a human keeping the register). The workbook is where they get READ, next to
    the module inventory they change the cost of, so they are copied in rather
    than regenerated. A missing file is not an error: both sheets are optional.

    A headers-only file is a register that was STAGED and is still empty —
    which is exactly how the ticket register is meant to start — so it gets
    its sheet, its CSV sibling and its `0 total` summary line. Only a file
    that is not there at all returns None.
    """
    if not path or not path.exists():
        return None
    rows = []
    with path.open(encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        missing = [c for c in columns if c not in (reader.fieldnames or [])]
        if missing:
            problems.append(f"{title}: {path.name} is missing column(s) "
                            f"{', '.join(missing)} — sheet skipped")
            return None
        for line_no, row in enumerate(reader, 2):
            if validate:
                for message in validate(row):
                    problems.append(f"{title}: {path.name}:{line_no} {message}")
            rows.append([row.get(c, "") for c in columns])
    sheet = workbook.create_sheet(title)
    sheet.append(columns)
    for row in rows:
        sheet.append(row)
    style_sheet(sheet, widths, "A2")
    return columns, rows


def validate_studio(row):
    classification = (row.get("classification") or "").strip()
    if classification not in STUDIO_CLASSIFICATIONS:
        yield (f"unknown classification {classification!r} "
               f"(expected one of {', '.join(sorted(STUDIO_CLASSIFICATIONS))})")
    if not (row.get("technical_name") or "").strip():
        yield "row has no technical_name"


def validate_ticket(row):
    status = (row.get("status") or "").strip()
    if status not in TICKET_STATUSES:
        yield (f"unknown status {status!r} "
               f"(expected one of {', '.join(sorted(TICKET_STATUSES))})")
    # A ticket without a token cannot be chased: the token is the only
    # identifier the upgrade client prints and the only one support can act on.
    if not (row.get("token") or "").strip():
        yield f"ticket {row.get('id') or '(unnumbered)'} has no upgrade token"
    # sop.md 7.2 makes an unresolved ticket a go-live blocker, and a waiver has
    # to point at the written waiver.
    if status == "waived" and not (row.get("resolution") or "").strip():
        yield (f"ticket {row.get('id') or '(unnumbered)'} is waived with no "
               f"resolution — record the written waiver")


def validate_area(area):
    """The inventory column answers to the same closed list as the
    requirements (AREA_ORDER), so an area copied from a module row into a
    requirement always passes. Compound values (`Sales / Reporting`) fail
    here on purpose: a module can produce requirements in two areas, and the
    bucket is assigned per REQUIREMENT, not per module."""
    if area not in AREA_ORDER:
        yield (f"Functional area {area!r} is not one of the "
               f"{len(AREA_ORDER)} buckets ({', '.join(AREA_ORDER)}) — "
               f"assign exactly one; secondary areas go in Module Purpose")


def validate_oca_repo(cell):
    """`none`, or `OCA/<repo>` (optionally `(series,…)`), `; `-joined."""
    if cell == OCA_NONE:
        return
    if cell == OCA_UNVERIFIED:
        yield ("OCA repo still carries the offline seed "
               f"{OCA_UNVERIFIED!r} — run oca_check.py before building")
        return
    if not cell:
        yield f"OCA repo is empty — write {OCA_NONE!r} when nothing matched"
        return
    for part in cell.split("; "):
        if not OCA_REPO_RE.match(part):
            yield (f"OCA repo part {part!r} is neither {OCA_NONE!r} nor "
                   f"OCA/<repo> — scan provenance goes to the Evidence sheet")


def validate_oca_alt(cell, column, oca_repo):
    """`none`, `already OCA: <repo>/<module>`, or `<repo>/<module>
    (full|partial)` — up to three, `; `-joined.

    A module that IS the OCA module is not an alternative to itself: when
    the OCA repo column names a verified upstream, `none` is false.

    The sentinel is NOT tolerated here: an unanswered column is reported by
    the sentinel count, which is what makes it exit non-zero.
    """
    if cell == OCA_NONE:
        if OCA_REPO_RE.match(oca_repo):
            yield (f"{column} is {OCA_NONE!r} but OCA repo names {oca_repo} "
                   f"— this module IS the OCA module: "
                   f"already OCA: <repo>/<module>")
        return
    for part in cell.split("; "):
        if not (OCA_ALREADY_RE.match(part) or OCA_ALT_RE.match(part)):
            yield (f"{column} part {part!r} is not {OCA_NONE!r}, "
                   f"'already OCA: <repo>/<module>', or "
                   f"'<repo>/<module> (full|partial)'")


def validate_native(verdict, column, notes_column):
    """The verdict column answers to a closed four-value list, so the sheet can
    be filtered on it and the summary can count it."""
    if verdict not in NATIVE_VERDICTS:
        yield (f"{column} is {verdict!r} — expected exactly one of "
               f"{', '.join(NATIVE_VERDICTS)}; the explanation goes in "
               f"{notes_column!r}")


def check_cap(label, value, cap):
    """One problem line for every documented length cap, in one shape."""
    if len(value) > cap:
        yield (f"{label} is {len(value)} chars (max {cap}): {value[:60]}…")


def split_native(entry, module, native_column, notes_column, review):
    """Verdict + description, from either shape of the phase-1 record.

    `native` is the verdict and `native_notes` the description. The legacy
    shape packed both into `native` as `<verdict>: <text>` inside 50
    characters, which is what made the description unreadable; it is still
    accepted — split on the first `: ` — and recorded as a review line, so
    the brief that produced it gets fixed rather than the data rejected.
    """
    verdict = (entry.get("native") or SENTINEL).strip()
    notes = (entry.get("native_notes") or "").strip()
    if verdict not in NATIVE_VERDICTS and ": " in verdict:
        head, _, tail = verdict.partition(": ")
        if head.strip() in NATIVE_VERDICTS:
            verdict, notes = head.strip(), notes or tail.strip()
            review.append(
                f"{module}: {native_column} carried the legacy "
                f"'<verdict>: <text>' cell — split into {verdict!r} + "
                f"{notes_column}; fix the phase-1 brief that wrote it")
    return verdict, notes


def build_inventory(workbook, workdir, source, native_column, notes_column,
                    oca_column, problems, review):
    enrich = merge_group_json(workdir, "enrich_*.json", problems)
    alternatives = merge_group_json(workdir, "oca_alt_*.json", problems)
    columns = BASE_COLUMNS + [native_column, notes_column, oca_column]
    # The columns the workbook presents as answers. A sentinel in any of them
    # is an unanswered question shipped as a verdict.
    delivered = ["Module Purpose", "3rd party app?", "Complexity/risk",
                 "Upgrade action", "Functional area", native_column,
                 notes_column, oca_column]
    sentinels = Counter()

    inventory_rows, evidence_rows, by_module = [], [], {}
    for seeded in csv.DictReader(source.open(encoding="utf-8-sig")):
        # Module Name is the technical name and nothing else — it is the key
        # the enrichment JSON, the requirement sources and the Traceability
        # sheet all join on, so it is used whole, never split.
        module = seeded["Module Name"].strip()
        enriched = enrich.get(module)
        alternative = alternatives.get(module)
        if not enriched:
            problems.append(f"{module}: no enrichment row (phase 1)")
        if not alternative:
            problems.append(f"{module}: no OCA alternative row (phase 2)")
        enriched, alternative = enriched or {}, alternative or {}

        verdict, native_notes = split_native(enriched, module, native_column,
                                             notes_column, review)
        native_evidence = (enriched.get("native_evidence") or "").strip()
        alt_evidence = (alternative.get("oca_alt_evidence") or "").strip()

        row = dict(seeded)
        row["Module Purpose"] = enriched.get("purpose") or row["Module Purpose"]
        row["3rd party app?"] = (enriched.get("third_party")
                                 or row["3rd party app?"])
        row["3rd party app link"] = (enriched.get("third_party_link")
                                     or row["3rd party app link"])
        row["Complexity/risk"] = (enriched.get("complexity")
                                  or row["Complexity/risk"])
        row["Upgrade action"] = (enriched.get("upgrade_action")
                                 or row["Upgrade action"])
        row["Functional area"] = (enriched.get("functional_area")
                                  or row["Functional area"])
        row[native_column] = verdict
        row[notes_column] = native_notes
        row[oca_column] = alternative.get("oca_alt", SENTINEL)

        for column in delivered:
            if (row.get(column) or "").strip() == SENTINEL:
                sentinels[column] += 1
                problems.append(f"{module}: {column} is {SENTINEL}")

        area = (row.get("Functional area") or "").strip()
        oca_repo = (row.get("OCA repo") or "").strip()
        oca_alt = (row.get(oca_column) or "").strip()
        # A sentinel cell has already been reported, once, by the count above;
        # running the vocabulary checks on it would say the same thing twice.
        checks = itertools.chain(
            validate_area(area) if area != SENTINEL else (),
            validate_oca_repo(oca_repo),
            validate_native(verdict, native_column, notes_column)
            if verdict != SENTINEL else (),
            check_cap(notes_column, native_notes, NATIVE_NOTES_MAX),
            check_cap("native? evidence", native_evidence, EVIDENCE_MAX),
            check_cap("OCA alternative evidence", alt_evidence, EVIDENCE_MAX),
            validate_oca_alt(oca_alt, oca_column, oca_repo)
            if oca_alt != SENTINEL else ())
        for message in checks:
            problems.append(f"{module}: {message}")

        inventory_rows.append([row.get(c, "") for c in columns])
        by_module[module] = row

        notes = " | ".join(x for x in [enriched.get("notes", ""),
                                       alternative.get("notes", "")] if x)
        evidence_rows.append([
            module, native_evidence, alt_evidence,
            enriched.get("vendor_release", ""),
            enriched.get("complexity_rationale", ""),
            enriched.get("upgrade_action_rationale", ""), notes])

    sheet = workbook.create_sheet("Module Inventory")
    sheet.append(columns)
    for row in inventory_rows:
        sheet.append(row)
    style_sheet(sheet, [28, 36, 70, 14, 8, 40, 18, 12, 34, 34, 12, 16, 16,
                        14, 60, 50], "B2")

    evidence = workbook.create_sheet("Inventory Evidence")
    evidence.append(EVIDENCE_COLUMNS)
    for row in evidence_rows:
        evidence.append(row)
    style_sheet(evidence, [30, 70, 70, 14, 44, 60, 60], "B2")

    return columns, inventory_rows, evidence_rows, by_module, sentinels


def mask_literals(text):
    """Blank out `backticked` and "quoted" spans before the wording rules run.

    Length is preserved so clause boundaries either side of the literal read
    the same to the patterns.
    """
    return QUOTED_LITERAL.sub(lambda m: "_" * len(m.group(0)), text)


def weak_wording(text):
    """Every weak-wording rule that fires, as (label, matched text).

    One line per rule, not per occurrence: the author fixes a class of wording,
    not a list of offsets.
    """
    masked = mask_literals(text)
    hits = []
    for label, pattern in WEAK_WORDS:
        match = pattern.search(masked)
        if match:
            hits.append((label, match.groupdict().get("w")
                         or match.group(0)))
    return hits


def normalise_requirement(entry, fallback_id, problems):
    """Make one entry safe to read with [] everywhere downstream.

    A missing key is already a problem line by the time this runs. The report
    is the only check that the fan-out agents followed their briefs, so it has
    to survive a malformed entry rather than be replaced by its traceback: fill
    what is missing, keep validating, exit non-zero at the end.

    `sources` is the one non-string key, and a one-module requirement reads
    naturally as a scalar — a bare string is iterable, so left alone it would
    emit one `unknown source` line per letter. Type it here, once.
    """
    entry.setdefault("tmp_id", fallback_id)
    sources = entry.get("sources", [])
    kind = type(sources).__name__
    if isinstance(sources, list):
        bad = next((m for m in sources if not isinstance(m, str)), None)
        if bad is not None:
            kind = f"list containing {type(bad).__name__}"
    if kind != "list":
        problems.append(f"{entry['tmp_id']}: sources must be a list of "
                        f"module names, got {kind}")
        sources = []
    entry["sources"] = sources
    for key in FR_KEYS - {"sources"}:
        entry.setdefault(key, "")


def group_of(tmp_id):
    """The agent group of a `tmp_id`: everything before the last `-`.

    `A-01` -> `A`, `SALES-07` -> `SALES`, `A1-02` -> `A1`. Comparing the first
    character alone made `G1`..`G6` — the obvious group naming — one single
    group, which silently disabled the only cross-group review the fan-out has.
    """
    return tmp_id.rsplit("-", 1)[0] if "-" in tmp_id else tmp_id


def load_requirements(workdir, inventory, handled_values, problems):
    requirements = []
    for path in sorted(workdir.glob("fr_*.json")):
        if path.name in {"fr_groups.json", "fr_merges.json", "fr_notes.json"}:
            continue
        for entry in json.loads(path.read_text()):
            tmp_id = entry.get("tmp_id", path.name)
            if set(entry) != FR_KEYS:
                problems.append(f"{tmp_id}: key mismatch "
                                f"{sorted(set(entry) ^ FR_KEYS)}")
            normalise_requirement(entry, tmp_id, problems)
            if entry.get("handled") not in handled_values:
                problems.append(f"{tmp_id}: handled={entry.get('handled')!r}")
            if entry.get("status_source") not in STATUSES:
                problems.append(
                    f"{tmp_id}: status={entry.get('status_source')!r}")
            if entry.get("functional_area") not in AREA_ORDER:
                problems.append(
                    f"{tmp_id}: area={entry.get('functional_area')!r}")
            text = entry.get("requirement", "")
            for label, word in weak_wording(text):
                problems.append(
                    f"{tmp_id}: weak wording [{label}] {word!r}: {text[:80]}")
            # Case-insensitive, and the message says so: a `Shall` capitalised
            # by an editor is the convention, not a missing keyword.
            if not re.search(r"\bshall\b", text, re.I):
                problems.append(f'{tmp_id}: no "shall" (case-insensitive) '
                                f'in: {text[:80]}')
            for module in entry["sources"]:
                if module not in inventory:
                    problems.append(f"{tmp_id}: unknown source {module}")
            requirements.append(entry)
    return requirements


def apply_reconciliation(workdir, requirements):
    by_tmp = {e.get("tmp_id", ""): e for e in requirements}
    merges_path = workdir / "fr_merges.json"
    if merges_path.is_file():
        for drop, spec in json.loads(merges_path.read_text()).items():
            dropped, kept = by_tmp.get(drop), by_tmp.get(spec.get("keep"))
            if not dropped or not kept:
                continue
            for module in dropped.get("sources", []):
                if module not in kept.setdefault("sources", []):
                    kept["sources"].append(module)
            kept["evidence"] = (kept.get("evidence", "") + " | "
                                + dropped.get("evidence", ""))[:400]
            kept["notes"] = " | ".join(
                x for x in [kept.get("notes", ""), spec.get("note", "")] if x)
            if spec.get("requirement"):
                kept["requirement"] = spec["requirement"]
            requirements.remove(dropped)
    notes_path = workdir / "fr_notes.json"
    if notes_path.is_file():
        for tmp_id, note in json.loads(notes_path.read_text()).items():
            entry = by_tmp.get(tmp_id)
            if entry:
                entry["notes"] = " | ".join(
                    x for x in [entry.get("notes", ""), note] if x)


def near_duplicates(requirements, threshold=0.33):
    stop = {"system", "shall", "when", "while", "user", "order", "where"}

    def tokens(text):
        return {w for w in re.findall(r"[a-z]+", text.lower())
                if len(w) > 3 and w not in stop}

    pairs = []
    for left, right in itertools.combinations(requirements, 2):
        if group_of(left.get("tmp_id", "")) == group_of(
                right.get("tmp_id", "")):
            continue
        a, b = tokens(left["requirement"]), tokens(right["requirement"])
        if a and b:
            score = len(a & b) / len(a | b)
            if score >= threshold:
                pairs.append((round(score, 2), left["id"], right["id"],
                              left["requirement"][:70],
                              right["requirement"][:70]))
    return sorted(pairs, reverse=True)


def build_requirements(workbook, workdir, inventory, native_column,
                       oca_column, handled_base, problems, review):
    handled_values = {handled_base, "OCA", "no"}
    requirements = load_requirements(workdir, inventory, handled_values,
                                     problems)
    if not requirements:
        return [], [], []

    apply_reconciliation(workdir, requirements)

    requirements.sort(key=lambda e: (
        AREA_ORDER.index(e.get("functional_area"))
        if e.get("functional_area") in AREA_ORDER else 99,
        e.get("tmp_id", "")))
    for number, entry in enumerate(requirements, 1):
        entry["id"] = f"FR-{number:03d}"

    columns = ["ID", "Requirement", "Type", "Functional area", "Actor",
               "Source customization(s)", "Status in source", "Handled?",
               "Handled by", "Handled notes", "Verification", "Evidence",
               "Notes", "Decision (Keep/Drop/Defer)"]
    rows = [[e["id"], e.get("requirement", ""), e.get("type", ""),
             e.get("functional_area", ""), e.get("actor", ""),
             "; ".join(e.get("sources", [])),
             e.get("status_source", "") + (f" — {e.get('status_note')}"
                                           if e.get("status_note") else ""),
             e.get("handled", ""), e.get("handled_by", ""),
             e.get("handled_notes", ""), e.get("verification", ""),
             e.get("evidence", ""), e.get("notes", ""), ""]
            for e in requirements]

    fills = {handled_base: PatternFill("solid", fgColor="E2EFDA"),
             "OCA": PatternFill("solid", fgColor="DDEBF7"),
             "no": PatternFill("solid", fgColor="FCE4D6")}
    sheet = workbook.create_sheet("Functional Requirements")
    sheet.append(columns)
    for entry, row in zip(requirements, rows):
        sheet.append(row)
        fill = fills.get(entry.get("handled"))
        if fill:
            sheet.cell(row=sheet.max_row, column=8).fill = fill
        if entry.get("status_source") != "active":
            sheet.cell(row=sheet.max_row, column=7).font = Font(
                color="C00000", bold=True)
    style_sheet(sheet, [9, 70, 16, 14, 16, 34, 18, 10, 34, 34, 44, 34, 36, 14],
                "C2")

    by_module = defaultdict(list)
    for entry in requirements:
        for module in entry.get("sources", []):
            by_module[module].append(entry)
    trace_columns = ["Module", "Requirement IDs", "# reqs",
                     f"# {handled_base}", "# OCA", "# no",
                     "Inventory: Upgrade action", f"Inventory: {native_column}",
                     f"Inventory: {oca_column}"]
    trace_rows = []
    for module in sorted(inventory, key=str.lower):
        entries = by_module.get(module, [])
        counts = Counter(e.get("handled") for e in entries)
        trace_rows.append([
            module, ", ".join(e["id"] for e in entries), len(entries),
            counts[handled_base], counts["OCA"], counts["no"],
            inventory[module]["Upgrade action"],
            inventory[module][native_column], inventory[module][oca_column]])
    trace = workbook.create_sheet("Traceability")
    trace.append(trace_columns)
    for row in trace_rows:
        trace.append(row)
    style_sheet(trace, [44, 40, 8, 11, 7, 6, 18, 44, 50], "B2")

    uncovered = sorted(set(inventory) - set(by_module))
    if uncovered:
        problems.append(f"modules with no requirement: {', '.join(uncovered)}")
    for score, left, right, left_text, right_text in near_duplicates(
            requirements):
        review.append(f"near-duplicate {score} {left}/{right}: "
                      f"{left_text} || {right_text}")

    return (columns, rows, requirements), (trace_columns, trace_rows), \
        requirements


def default_if_present(workdir, name):
    """Optional inputs are picked up when they exist, never demanded."""
    candidate = workdir / name
    return candidate if candidate.exists() else None


def main():
    parser = argparse.ArgumentParser(
        description="Build the upgrade-preparation workbook.")
    parser.add_argument("--workdir", type=Path, required=True,
                        help="Dir holding inventory.csv and the agent JSON")
    parser.add_argument("--inventory", type=Path,
                        help="Seeded inventory CSV "
                             "(default: <workdir>/inventory.csv)")
    parser.add_argument("--target-version", default="19.0",
                        help="Target series, e.g. 19.0 (default: 19.0)")
    parser.add_argument("-o", "--output", type=Path, required=True,
                        help="Output xlsx path")
    parser.add_argument("--studio", type=Path,
                        help="studio_inventory.py CSV "
                             "(default: <workdir>/studio.csv if present)")
    parser.add_argument("--tickets", type=Path,
                        help="Odoo support ticket register CSV "
                             "(default: <workdir>/tickets.csv if present)")
    args = parser.parse_args()

    major = args.target_version.split(".")[0]
    native_column = f"{major} native?"
    notes_column = f"{major} native notes"
    oca_column = f"OCA {major} alternative"
    handled_base = f"base {major}"
    source = args.inventory or args.workdir / "inventory.csv"
    problems, review = [], []

    workbook = Workbook()
    workbook.remove(workbook.active)
    inv_columns, inv_rows, evidence_rows, inventory, sentinels = \
        build_inventory(workbook, args.workdir, source, native_column,
                        notes_column, oca_column, problems, review)
    fr, trace, requirements = build_requirements(
        workbook, args.workdir, inventory, native_column, oca_column,
        handled_base, problems, review)

    studio = build_passthrough_sheet(
        workbook, args.studio or default_if_present(args.workdir, "studio.csv"),
        "Studio", STUDIO_COLUMNS, [14, 34, 24, 30, 20, 8, 18, 40],
        problems, validate_studio)
    tickets = build_passthrough_sheet(
        workbook, args.tickets or default_if_present(args.workdir, "tickets.csv"),
        "Tickets", TICKET_COLUMNS, [10, 12, 52, 26, 14, 18, 40, 34],
        problems, validate_ticket)

    # Requirements sheets are appended after the inventory ones; the workbook
    # reads inventory-first, so reorder to put requirements second.
    order = ["Module Inventory", "Functional Requirements", "Traceability",
             "Inventory Evidence", "Studio", "Tickets"]
    workbook._sheets.sort(
        key=lambda ws: order.index(ws.title) if ws.title in order else 99)
    workbook.save(args.output)

    stem = args.output.with_suffix("")
    write_csv(Path(f"{stem}_inventory.csv"), inv_columns, inv_rows)
    write_csv(Path(f"{stem}_evidence.csv"), EVIDENCE_COLUMNS, evidence_rows)
    if requirements:
        write_csv(Path(f"{stem}_requirements.csv"), fr[0], fr[1])
        write_csv(Path(f"{stem}_traceability.csv"), trace[0], trace[1])
    if studio:
        write_csv(Path(f"{stem}_studio.csv"), studio[0], studio[1])
    if tickets:
        write_csv(Path(f"{stem}_tickets.csv"), tickets[0], tickets[1])

    print(f"saved {args.output}: "
          f"{[(ws.title, ws.max_row - 1) for ws in workbook]}")
    if studio:
        print("studio:", dict(Counter(r[STUDIO_COLUMNS.index("classification")]
                                      for r in studio[1])))
    if tickets:
        open_tickets = [r for r in tickets[1]
                        if r[TICKET_COLUMNS.index("status")] not in ("resolved", "waived")]
        # sop.md 7.2: every ticket resolved or waived before a production run.
        print(f"tickets: {len(tickets[1])} total, {len(open_tickets)} still blocking")
    if requirements:
        print("handled:", dict(Counter(e["handled"] for e in requirements)))
        print("status:", dict(Counter(e["status_source"] for e in requirements)))
        print("area:", dict(Counter(e["functional_area"]
                                    for e in requirements)))
    # Per-column sentinel counts: the difference between "schema-valid" and
    # "answered". Printed even when empty, so a clean run says so.
    print("sentinels:", dict(sentinels))
    print(f"problems: {len(problems)}")
    for problem in problems[:80]:
        print("  ", problem)
    if len(problems) > 80:
        print(f"   ... {len(problems) - 80} more")
    print(f"to review (not blocking): {len(review)}")
    for item in review[:40]:
        print("  ", item)
    if len(review) > 40:
        print(f"   ... {len(review) - 40} more")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
